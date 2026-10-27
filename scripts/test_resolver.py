import openpyxl
import os
import re
from dotenv import load_dotenv
from supabase import create_client

load_dotenv()
supabase = create_client(os.environ['SUPABASE_URL'], os.environ['SUPABASE_KEY'])

locations = supabase.table('locations').select('*').execute().data
print(f"Loaded {len(locations)} existing locations from Supabase.")

# Build robust lookup dictionaries
# 1. By exact store_number
loc_by_num = {loc['store_number']: loc for loc in locations}

# 2. By clean normalized name
def normalize_name(s):
    if not s: return ""
    s = s.lower()
    s = re.sub(r'\(.*?\)', '', s) # remove parenthetical notes like (Oct 17)
    s = s.replace('mobile klinik', '').replace('frchs-', '').replace('franchise', '').strip()
    s = re.sub(r'[^a-z0-9]', '', s) # remove spaces, hyphens, punctuation
    return s

loc_by_name = {}
for loc in locations:
    norm_n = normalize_name(loc['name'])
    norm_sn = normalize_name(loc.get('short_name') or '')
    if norm_n:
        loc_by_name[norm_n] = loc
    if norm_sn:
        loc_by_name[norm_sn] = loc

# Special alias mapping if needed
aliases = {
    'midislandnanaimo': 'nanaimowest',
    'nanaimo': 'midislandnanaimo',
    'kp': 'kildonanplace',
    'wem': 'westedmontonmall',
    'lansdowne': 'landsdowne',
    'landsdowne': 'landsdowne'
}

wb = openpyxl.load_workbook('Updated Franchise Scorecard - Oct MTD EOD - 10-25-26.xlsx', read_only=True, data_only=True)

# Test resolving Lethbridge in all sheets
print("\n--- RESOLVING LETHBRIDGE ACROSS ALL SHEETS ---")
for sheet in wb.sheetnames:
    if 'koodo' in sheet.lower() or 'outdoor' in sheet.lower():
        continue
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    for r_idx, r in enumerate(rows[:100], start=1):
        # check if lethbridge in row
        r_str = [str(c or '').lower() for c in r[:6]]
        if any('lethbridge' in c for c in r_str):
            # How would our new resolver resolve this row?
            raw_code = r[0] if len(r)>0 else None
            raw_name = r[1] if len(r)>1 else None
            
            # 1. Try matching by Name first
            matched_by = None
            resolved_loc = None
            
            norm_target = normalize_name(str(raw_name or ''))
            if norm_target in loc_by_name:
                resolved_loc = loc_by_name[norm_target]
                matched_by = f"NAME ({raw_name} -> {resolved_loc['name']} #{resolved_loc['store_number']})"
            else:
                # search row for store name
                for c in r[:6]:
                    norm_c = normalize_name(str(c or ''))
                    if norm_c in loc_by_name:
                        resolved_loc = loc_by_name[norm_c]
                        matched_by = f"ROW_CELL_NAME ({c} -> {resolved_loc['name']} #{resolved_loc['store_number']})"
                        break
                        
            print(f"Sheet '{sheet:12s}' Row {r_idx:2d}: RawCol0={str(raw_code):8s} | RawCol1={str(raw_name):12s} | Resolved: {matched_by}")
