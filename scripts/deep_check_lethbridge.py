import openpyxl
import os
from dotenv import load_dotenv
from supabase import create_client

load_dotenv()
supabase = create_client(os.environ['SUPABASE_URL'], os.environ['SUPABASE_KEY'])

# Get Lethbridge location id
loc_326 = supabase.table('locations').select('*').eq('store_number', 326).execute().data[0]
loc_id = loc_326['id']

# Get periods
periods_list = supabase.table('scorecard_periods').select('*').execute().data
period_by_label = {p['label'].strip(): p['id'] for p in periods_list}
period_label_by_id = {p['id']: p['label'].strip() for p in periods_list}

# Get current DB kpis for Lethbridge
db_kpis = {k['period_id']: k for k in supabase.table('monthly_kpis').select('*').eq('location_id', loc_id).execute().data}

wb = openpyxl.load_workbook('Copy of Franchise Scorecard - Oct MTD EOD - 10-25-26.xlsx', read_only=True, data_only=True)

print(f"Checking Lethbridge (Store 326, DB ID: {loc_id}):")

sheets = ['Dec 2025 ', 'Nov 25', 'Oct 2025', 'Sept 2025', 'Jan 26', 'Feb 26', 'Mar 26', 'April 26', 'May 26', 'June 26', 'July 26', 'Aug 26', 'Sept 26', 'Oct 26 MTD']

for s_name in sheets:
    if s_name not in wb.sheetnames:
        continue
    ws = wb[s_name]
    rows = list(ws.iter_rows(values_only=True))
    
    # find row for lethbridge
    target_row = None
    target_row_idx = None
    for r_idx, r in enumerate(rows):
        if r_idx > 100: break
        for c in r[:6]:
            if c is not None and str(c).strip().lower() == 'lethbridge':
                target_row = r
                target_row_idx = r_idx + 1
                break
        if target_row: break
        
    p_label = s_name.strip()
    pid = period_by_label.get(p_label)
    db_data = db_kpis.get(pid) if pid else None
    
    print(f"\n==========================================")
    print(f"Sheet: '{s_name}' (Excel Row {target_row_idx}) | Period: '{p_label}'")
    print(f"==========================================")
    if not target_row:
        print("  NO LETHBRIDGE ROW FOUND IN SHEET!")
        continue
        
    # Inspect Excel headers in row 4 and row 8
    header_cat = rows[3] if len(rows) > 3 else []
    header_sub = rows[7] if len(rows) > 7 else []
    
    # Let's extract key values from target_row
    # Look for gross revenue, traffic, act, cpo, rep
    excel_vals = {}
    for c_idx, val in enumerate(target_row):
        cat = str(header_cat[c_idx] or '').strip().replace('\n', ' ') if c_idx < len(header_cat) else ''
        sub = str(header_sub[c_idx] or '').strip().replace('\n', ' ') if c_idx < len(header_sub) else ''
        sub_l = sub.lower()
        cat_l = cat.lower()
        
        # Revenue
        if 'total gross revenue' in sub_l and 'yoy' not in sub_l and 'mom' not in sub_l:
            excel_vals['rev'] = val
            excel_vals['rev_col'] = (c_idx, sub)
        elif 'all revenue' in sub_l and 'rev' not in excel_vals:
            excel_vals['rev'] = val
            excel_vals['rev_col'] = (c_idx, sub)
            
        # Traffic
        if 'traffic' in sub_l and ('total' in sub_l or 'mtd' in sub_l) and 'mom' not in sub_l and 'yoy' not in sub_l:
            if 'traffic' not in excel_vals:
                excel_vals['traffic'] = val
                excel_vals['traffic_col'] = (c_idx, sub)
        elif 'traffic' in cat_l and ('total' in sub_l or 'mtd' in sub_l or sub == ''):
            if 'traffic' not in excel_vals:
                excel_vals['traffic'] = val
                excel_vals['traffic_col'] = (c_idx, sub)
                
        # Activations
        if 'activation' in sub_l and ('total' in sub_l or 'all' in sub_l) and ('trans' in sub_l or 'mtd' in sub_l) and 'mom' not in sub_l and 'yoy' not in sub_l and 'attached' not in sub_l and 'attach' not in sub_l:
            if 'act' not in excel_vals:
                excel_vals['act'] = val
                excel_vals['act_col'] = (c_idx, sub)
                
        # CPO units
        if 'cpo units' in sub_l and ('all' in sub_l or 'total' in sub_l or 'mtd' in sub_l) and 'mom' not in sub_l and 'yoy' not in sub_l:
            if 'cpo' not in excel_vals:
                excel_vals['cpo'] = val
                excel_vals['cpo_col'] = (c_idx, sub)
                
        # Repairs
        if 'total oow repairs' in sub_l or ('oow' in sub_l and 'repairs' in sub_l and 'mom' not in sub_l and 'yoy' not in sub_l and 'day' not in sub_l and 'rev' not in sub_l):
            if 'rep' not in excel_vals:
                excel_vals['rep'] = val
                excel_vals['rep_col'] = (c_idx, sub)

    if not db_data:
        print("  STATUS: NO RECORD IN DB FOR THIS PERIOD!")
    else:
        db_rev = db_data.get('gross_revenue_actual')
        db_traf = db_data.get('foot_traffic_total')
        db_act = db_data.get('activations_total')
        db_cpo = db_data.get('cpo_units_actual')
        db_rep = db_data.get('repairs_oow_volume')
        
        ex_rev = excel_vals.get('rev')
        ex_traf = excel_vals.get('traffic')
        ex_act = excel_vals.get('act')
        ex_cpo = excel_vals.get('cpo')
        ex_rep = excel_vals.get('rep')
        
        rev_match = (db_rev == ex_rev) or (abs((db_rev or 0) - (ex_rev or 0)) < 0.01 if isinstance(ex_rev, (int, float)) and isinstance(db_rev, (int, float)) else False)
        traf_match = (db_traf == ex_traf)
        act_match = (db_act == ex_act)
        cpo_match = (db_cpo == ex_cpo)
        rep_match = (db_rep == ex_rep)
        
        all_match = rev_match and traf_match and act_match and cpo_match and rep_match
        
        print(f"  Match Status: {'MATCH' if all_match else 'DISCREPANCY DETECTED!'}")
        print(f"    Revenue:  Excel={ex_rev}  |  DB={db_rev}  -> {'OK' if rev_match else 'MISMATCH'}")
        print(f"    Traffic:  Excel={ex_traf}  |  DB={db_traf}  -> {'OK' if traf_match else 'MISMATCH'}")
        print(f"    Act:      Excel={ex_act}  |  DB={db_act}  -> {'OK' if act_match else 'MISMATCH'}")
        print(f"    CPO:      Excel={ex_cpo}  |  DB={db_cpo}  -> {'OK' if cpo_match else 'MISMATCH'}")
        print(f"    Repairs:  Excel={ex_rep}  |  DB={db_rep}  -> {'OK' if rep_match else 'MISMATCH'}")
